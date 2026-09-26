import io

import openpyxl
import pandas as pd
import pytest

import etf_analyst_report_complete as pipeline


class _FakeResponse:
    def __init__(self, content: bytes, status_code=200):
        self.content = content
        self.status_code = status_code
        self.text = content.decode("utf-8", errors="ignore")
        self.headers = {"content-type": "text/csv"}


def _make_oef_csv():
    rows = ["Ticker,Name,Weight (%)"]
    for i in range(100):
        ticker = f"T{i:03d}"
        weight = 2.0 if i == 0 else 98.0 / 99.0
        rows.append(f"{ticker},Company {i},{weight:.6f}")
    return ("\n".join(rows) + "\n").encode("utf-8")


def test_oef_is_in_required_etf_config():
    assert "OEF" in pipeline.ETFS
    assert pipeline.ETF_CONFIG["OEF"]["product_id"] == "239723"
    assert pipeline.ETF_CONFIG["OEF"]["file_name"] == "OEF_holdings"
    assert pipeline.ETF_CONFIG["OEF"]["min_rows"] >= 90


def test_oef_blackrock_csv_parser_uses_official_product_id(monkeypatch):
    seen_urls = []

    def fake_get(url, headers=None, timeout=None):
        seen_urls.append(url)
        return _FakeResponse(_make_oef_csv())

    monkeypatch.setattr(pipeline.requests, "get", fake_get)

    raw = pipeline.pull_blackrock_oef()
    normalized = pipeline.normalize_holdings(raw, "OEF")

    assert len(normalized) == 100
    assert normalized["Weight"].sum() == pytest.approx(100.0, abs=0.001)
    assert seen_urls
    assert "239723" in seen_urls[0]
    assert seen_urls[0].endswith("/latest-holdings.csv")


def test_forward_pe_charts_live_on_dedicated_tab(tmp_path):
    output = tmp_path / "layout.xlsx"
    hist = pd.DataFrame(
        [
            {"Date": "2026-08-01", "ETF": "QQQ", "PE Ratio": 30.0, "Forward PE": 25.0},
            {"Date": "2026-09-01", "ETF": "QQQ", "PE Ratio": 31.0, "Forward PE": 26.0},
            {"Date": "2026-08-01", "ETF": "OEF", "PE Ratio": 28.0, "Forward PE": 23.0},
            {"Date": "2026-09-01", "ETF": "OEF", "PE Ratio": 29.0, "Forward PE": 24.0},
        ]
    )

    with pd.ExcelWriter(output, engine="xlsxwriter") as writer:
        workbook = writer.book
        summary = workbook.add_worksheet("Summary")
        writer.sheets["Summary"] = summary
        pipeline._write_summary_pe_history_data_and_charts(
            workbook,
            summary,
            writer,
            hist,
            chart_start_row=2,
            start_col=1,
        )

    wb = openpyxl.load_workbook(output, read_only=False, data_only=False)
    assert "Summary" in wb.sheetnames
    assert "Forward PE" in wb.sheetnames
    assert "PE_History_Data" in wb.sheetnames
    assert wb["PE_History_Data"].sheet_state == "hidden"
    assert len(wb["Summary"]._charts) == 2
    assert len(wb["Forward PE"]._charts) == 2
    assert wb["Forward PE"]["A1"].value == "Forward PE Dashboard"
    assert wb["Forward PE"]["B10"].value == "ETF"
    assert wb["Forward PE"]["G10"].value == "Premium/(Discount)"
    assert wb["Forward PE"]["H10"].value == "Percentile"
    assert wb["Forward PE"]["J10"].value == "Obs."
    assert wb["Forward PE"]["K10"].value == "Status"


def test_forward_pe_dashboard_stats_are_descriptive_and_stable():
    hist = pd.DataFrame(
        [
            {"Date": "2026-07-01", "ETF": "QQQ", "Forward PE": 20.0, "Forward PE coverage": 0.95},
            {"Date": "2026-08-01", "ETF": "QQQ", "Forward PE": 22.0, "Forward PE coverage": 0.96},
            {"Date": "2026-09-01", "ETF": "QQQ", "Forward PE": 24.0, "Forward PE coverage": 0.97},
            {"Date": "2026-07-01", "ETF": "OEF", "Forward PE": 20.0, "Forward PE coverage": 0.99},
            {"Date": "2026-08-01", "ETF": "OEF", "Forward PE": 20.0, "Forward PE coverage": 0.99},
            {"Date": "2026-09-01", "ETF": "OEF", "Forward PE": 19.0, "Forward PE coverage": 0.99},
        ]
    )
    hist["Date"] = pd.to_datetime(hist["Date"])
    stats = pipeline._build_forward_pe_dashboard_stats(hist).set_index("ETF")

    assert stats.loc["QQQ", "Current Forward PE"] == pytest.approx(24.0)
    assert stats.loc["QQQ", "1Y Avg"] == pytest.approx(22.0)
    assert stats.loc["QQQ", "Premium/(Discount)"] == pytest.approx(24.0 / 22.0 - 1.0)
    assert stats.loc["QQQ", "History Percentile"] == pytest.approx(1.0)
    assert stats.loc["QQQ", "Forward PE Coverage"] == pytest.approx(0.97)
    assert stats.loc["QQQ", "Status"] == "Above 1Y Avg"

    assert stats.loc["OEF", "Current Forward PE"] == pytest.approx(19.0)
    assert stats.loc["OEF", "Status"] == "Below 1Y Avg"
